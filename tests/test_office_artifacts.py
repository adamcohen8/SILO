from io import BytesIO
import json
import os
from pathlib import Path
import tempfile
import unittest
import zipfile

from local_ai_appliance.agent import DocumentAgent,Source
from local_ai_appliance.backend import PortableSession
from local_ai_appliance.contracts import Project,WorkspaceKind
from local_ai_appliance.document_tools import analyze_table,search_folder
from local_ai_appliance.documents import parse_document
from local_ai_appliance.folders import AccessDenied
from local_ai_appliance.model import ModelReply,ToolCall
from local_ai_appliance.office_artifacts import build_artifact
from local_ai_appliance.workspace import Workspace

TABLE={'title':'Capacity study','headers':['Asset','Capacity kg'],'rows':[['One',48],['Two',60],['Literal','=1+1']]}
DOC={'title':'Capacity study','sections':[{'heading':'Findings','paragraphs':['Available capacity is 108 kg.'],
                                          'table':{'headers':TABLE['headers'],'rows':TABLE['rows'][:2]}}]}
SLIDES={'title':'Capacity study','slides':[{'title':'Available capacity','bullets':['One: 48 kg','Two: 60 kg','Total: 108 kg']} ]}

class OfficeTests(unittest.TestCase):
    def test_office_files_embed_verified_references_and_source_digests(self):
        from docx import Document
        from openpyxl import load_workbook
        from pptx import Presentation
        source=Source('S1','folder','data.csv','a'*64,1,3,'rows','row')
        doc={'title':'Summary','sections':[{'heading':'Capacity','paragraphs':['108 kg. [source:S1]']}]}
        word=Document(BytesIO(build_artifact('docx',json.dumps(doc),sources=[source]).data))
        paragraphs='\n'.join(p.text for p in word.paragraphs)
        self.assertIn('[data.csv, row unit 1-3]',paragraphs)
        self.assertNotIn('[source:',paragraphs)
        self.assertIn(source.sha256,paragraphs)
        excel=load_workbook(BytesIO(build_artifact('xlsx',json.dumps(TABLE),sources=[source]).data))
        self.assertIn(source.sha256,excel['Sources']['A1'].value)
        slides=Presentation(BytesIO(build_artifact('pptx',json.dumps(SLIDES),sources=[source]).data))
        self.assertEqual(len(slides.slides),2)
        text='\n'.join(s.text for slide in slides.slides for s in slide.shapes if s.has_text_frame)
        self.assertIn(source.sha256,text)
        with self.assertRaises(AccessDenied):build_artifact('docx',json.dumps(doc))

    def test_deterministic_artifacts_roundtrip_in_separate_parser(self):
        for kind,spec in [('docx',DOC),('xlsx',TABLE),('csv',TABLE),('pptx',SLIDES)]:
            with self.subTest(kind=kind):
                artifact=build_artifact(kind,json.dumps(spec))
                self.assertEqual(artifact.data,build_artifact(kind,json.dumps(spec)).data)
                document=parse_document('opaque','report.'+kind,artifact.data)
                self.assertIn('48','\n'.join(document.units))
                self.assertIn('60','\n'.join(document.units))
                self.assertEqual(document.unit_kind,{'docx':'paragraph','xlsx':'row','csv':'row','pptx':'slide'}[kind])

    def test_word_independent_reader_styles_table_and_margins(self):
        from docx import Document
        document=Document(BytesIO(build_artifact('docx',json.dumps(DOC)).data))
        self.assertEqual(document.paragraphs[0].text,'Capacity study')
        self.assertEqual(document.paragraphs[1].style.name,'Heading1')
        self.assertEqual(document.tables[0].cell(2,1).text,'60')
        self.assertEqual(document.sections[0].left_margin.inches,1.0)

    def test_excel_independent_reader_values_literal_formula_frozen_header(self):
        from openpyxl import load_workbook
        workbook=load_workbook(BytesIO(build_artifact('xlsx',json.dumps(TABLE)).data),data_only=False)
        sheet=workbook.active
        self.assertEqual(sheet['B2'].value,48);self.assertEqual(sheet['B3'].value,60)
        self.assertEqual(sheet['B4'].value,'=1+1');self.assertEqual(sheet['B4'].data_type,'s')
        self.assertEqual(sheet.freeze_panes,'A2');self.assertTrue(sheet['A1'].font.bold)

    def test_powerpoint_independent_reader_slides_text_and_geometry(self):
        from pptx import Presentation
        presentation=Presentation(BytesIO(build_artifact('pptx',json.dumps(SLIDES)).data))
        self.assertEqual(len(presentation.slides),1)
        self.assertIn('108 kg','\n'.join(shape.text for shape in presentation.slides[0].shapes if shape.has_text_frame))
        self.assertEqual(presentation.slide_width,12192000)
        for shape in presentation.slides[0].shapes:
            self.assertLessEqual(shape.top+shape.height,presentation.slide_height)
            self.assertLessEqual(shape.left+shape.width,presentation.slide_width)

    def test_csv_formula_injection_escaped_and_numeric_negative_preserved(self):
        import csv,io
        spec={'title':'Literal','headers':['Value'],'rows':[['=2+2'],['  @cmd'],[-2],['-2']]}
        rows=list(csv.reader(io.StringIO(build_artifact('csv',json.dumps(spec)).data.decode())))
        self.assertEqual(rows[1][0],"'=2+2");self.assertEqual(rows[2][0],"'  @cmd")
        self.assertEqual(rows[3][0],'-2');self.assertEqual(rows[4][0],"'-2")

    def test_invalid_structure_controls_nan_and_formula_inputs_fail_closed(self):
        for kind,spec in [('docx',{'title':'x','sections':[]}),('xlsx',{'title':'x','headers':['a','a'],'rows':[[1,2]]}),
                          ('pptx',{'title':'x','slides':[{'title':'x','bullets':[]}]})]:
            with self.assertRaises(AccessDenied):build_artifact(kind,json.dumps(spec))
        for cell in ('bad\0',float('nan'),float('inf'),1e99):
            with self.assertRaises((AccessDenied,ValueError)):build_artifact('xlsx',json.dumps({'title':'x','headers':['a'],'rows':[[cell]]}))
        artifact=build_artifact('xlsx',json.dumps(TABLE));target=BytesIO()
        with zipfile.ZipFile(BytesIO(artifact.data)) as source,zipfile.ZipFile(target,'w') as out:
            for item in source.infolist():
                data=source.read(item)
                if item.filename=='xl/worksheets/sheet1.xml':data=data.replace(b'<v>48</v>',b'<f>1+1</f><v>48</v>')
                out.writestr(item,data)
        with self.assertRaises(AccessDenied):parse_document('f','formula.xlsx',target.getvalue())

@unittest.skipUnless(os.name=='posix','Portable office broker fixture')
class OfficeWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name).resolve();self.source=self.root/'source';self.dest=self.root/'dest'
        self.source.mkdir();self.dest.mkdir();self.workspace=Workspace();self.addCleanup(self.workspace.close)
        self.ro=self.workspace.grant(self.source,'read-only');self.rw=self.workspace.grant(self.dest,'editable')
        (self.source/'data.csv').write_text('Asset,Capacity kg\nOne,48\nTwo,60\nMissing,\nBad,unknown\n')

    def test_table_exact_analysis_counts_and_sheet_requirements(self):
        document,start,end,value=analyze_table(self.workspace,self.ro.id,'data.csv','Capacity kg')
        stats=value['statistics'];self.assertEqual(stats['sum'],'108');self.assertEqual(stats['mean'],'54')
        self.assertEqual((stats['numeric_count'],stats['missing_count'],stats['rejected_count']),(2,1,1))
        self.assertEqual((start,end),(1,5));self.assertEqual(document.unit_kind,'row')
        with self.assertRaises(AccessDenied):analyze_table(self.workspace,self.ro.id,'data.csv','Missing')

    def test_text_tool_cannot_bypass_structured_office_writers(self):
        for kind in ('docx','xlsx','csv','pptx'):
            with self.assertRaises(AccessDenied):self.workspace.propose(self.rw.id,'unsafe.'+kind,'=RUN()')
        self.assertFalse(list(self.dest.iterdir()))

    def test_cross_document_search_read_current_source_and_bound_partial_search(self):
        child=self.source/'nested';child.mkdir();(child/'find.md').write_text('The fictional radio bandwidth is 20 MHz.')
        document,kind,number,excerpt,value=search_folder(self.workspace,self.ro.id,'bandwidth')
        self.assertEqual(document.path,'nested/find.md');self.assertIn('20 MHz',excerpt)
        (child/'find.md').write_text('Changed current source')
        self.assertFalse(search_folder(self.workspace,self.ro.id,'bandwidth')[-1]['found'])
        for i in range(20):(self.source/f'file-{i}.txt').write_text('No matching keyword')
        result=search_folder(self.workspace,self.ro.id,'absent')[-1]
        self.assertFalse(result['search_complete']);self.assertEqual(result['searched_documents'],16)

    def test_artifact_pending_exact_review_save_no_overwrite_and_source_drift(self):
        document=self.workspace.open_document(self.ro.id,'data.csv')
        source=Source('S1',self.ro.id,'data.csv',document.sha256,1,5,'rows','row')
        proposal=self.workspace.propose_artifact(self.rw.id,'report.docx',json.dumps(DOC),[source])
        model=type('Model',(),{'complete':lambda self,m,t:ModelReply('Fixture')})()
        backend=PortableSession(Project('fixture','Fixture'),self.workspace,model)
        with self.assertRaises(AccessDenied):backend.approve(proposal.id)
        backend.review(proposal.id);self.assertFalse((self.dest/'report.docx').exists())
        backend.approve(proposal.id);self.assertEqual((self.dest/'report.docx').read_bytes(),proposal.data)
        with self.assertRaises(AccessDenied):backend.approve(proposal.id)
        with self.assertRaises(AccessDenied):self.workspace.propose_artifact(self.rw.id,'report.docx',json.dumps(DOC))
        stale=self.workspace.propose_artifact(self.rw.id,'stale.xlsx',json.dumps(TABLE),[source]);backend.review(stale.id)
        (self.source/'data.csv').write_text('Changed')
        with self.assertRaises(AccessDenied):backend.approve(stale.id)
        self.assertFalse((self.dest/'stale.xlsx').exists())

    def test_graph_analyze_then_propose_word_end_to_end(self):
        replies=iter([ModelReply(calls=(ToolCall('read','analyze_table',{'folder_id':self.ro.id,'path':'data.csv','column':'Capacity kg'}),)),
                      ModelReply(calls=(ToolCall('draft','propose_artifact',{'destination_folder':self.rw.id,'destination_path':'graph.docx',
                        'content':json.dumps({'title':'Summary','sections':[{'heading':'Capacity','paragraphs':['108 kg. [source:S1]']}]})}),)),
                      ModelReply('Draft ready. [source:S1]')])
        model=type('Model',(),{'complete':lambda self,m,t:next(replies)})()
        backend=PortableSession(Project('graph','Graph'),self.workspace,model)
        result=backend.ask('Analyze the capacity and prepare a Word report')
        self.assertEqual(result.citations[0].unit_kind,'row')
        self.assertEqual(len(result.proposals),1);proposal=result.proposals[0]
        backend.review(proposal.id);backend.approve(proposal.id)
        self.assertIn('108 kg','\n'.join(parse_document(self.rw.id,'graph.docx',proposal.data).units))
